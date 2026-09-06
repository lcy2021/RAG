"""Extract plain text from uploads for ingest loaders.

Routing is by *strategy* (how text is produced):

- text   — plain / structured text (txt, md, yaml, json, …)
- markup — HTML visible text
- layout — digital PDF / DOCX / PPTX text layer
- table  — CSV / XLSX (and DOCX tables) with header-aware merge
- ocr    — images and scanned PDFs (optional pytesseract)

``auto`` picks a strategy from extension/MIME (and PDF text density).
Specialized loaders force one strategy for A/B tests.
"""

from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass, field
from enum import StrEnum
from pathlib import Path

_TEXT_ENCODINGS = ("utf-8-sig", "utf-8", "gb18030", "latin-1")

_TEXT_EXT = {
    ".txt",
    ".md",
    ".markdown",
    ".rst",
    ".log",
    ".yml",
    ".yaml",
    ".xml",
    ".ini",
    ".cfg",
    ".conf",
}

_IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".tif", ".tiff", ".bmp"}

# Sparse digital PDF → treat as scan candidate (non-whitespace chars / page).
_PDF_SPARSE_CHARS_PER_PAGE = 40


class DocKind(StrEnum):
    TEXT = "text"
    HTML = "html"
    JSON = "json"
    CSV = "csv"
    PDF = "pdf"
    DOCX = "docx"
    XLSX = "xlsx"
    PPTX = "pptx"
    IMAGE = "image"


class LoaderStrategy(StrEnum):
    TEXT = "text"
    MARKUP = "markup"
    LAYOUT = "layout"
    TABLE = "table"
    OCR = "ocr"


_EXT_KIND: dict[str, DocKind] = {
    **{ext: DocKind.TEXT for ext in _TEXT_EXT},
    **{ext: DocKind.IMAGE for ext in _IMAGE_EXT},
    ".html": DocKind.HTML,
    ".htm": DocKind.HTML,
    ".json": DocKind.JSON,
    ".csv": DocKind.CSV,
    ".tsv": DocKind.CSV,
    ".pdf": DocKind.PDF,
    ".docx": DocKind.DOCX,
    ".xlsx": DocKind.XLSX,
    ".pptx": DocKind.PPTX,
}

_MIME_KIND: dict[str, DocKind] = {
    "text/plain": DocKind.TEXT,
    "text/markdown": DocKind.TEXT,
    "text/x-markdown": DocKind.TEXT,
    "text/html": DocKind.HTML,
    "application/json": DocKind.JSON,
    "text/csv": DocKind.CSV,
    "text/tab-separated-values": DocKind.CSV,
    "application/pdf": DocKind.PDF,
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": DocKind.DOCX,
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": DocKind.XLSX,
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": DocKind.PPTX,
    "image/png": DocKind.IMAGE,
    "image/jpeg": DocKind.IMAGE,
    "image/gif": DocKind.IMAGE,
    "image/webp": DocKind.IMAGE,
    "image/tiff": DocKind.IMAGE,
    "image/bmp": DocKind.IMAGE,
}

_LEGACY_HINTS = {
    ".doc": "legacy .doc is not supported; save as .docx",
    ".xls": "legacy .xls is not supported; save as .xlsx",
    ".ppt": "legacy .ppt is not supported; save as .pptx",
}

_KIND_STRATEGIES: dict[DocKind, frozenset[LoaderStrategy]] = {
    DocKind.TEXT: frozenset({LoaderStrategy.TEXT}),
    DocKind.JSON: frozenset({LoaderStrategy.TEXT}),
    DocKind.HTML: frozenset({LoaderStrategy.MARKUP}),
    DocKind.CSV: frozenset({LoaderStrategy.TABLE}),
    DocKind.XLSX: frozenset({LoaderStrategy.TABLE}),
    DocKind.DOCX: frozenset({LoaderStrategy.LAYOUT, LoaderStrategy.TABLE}),
    DocKind.PPTX: frozenset({LoaderStrategy.LAYOUT}),
    DocKind.PDF: frozenset({LoaderStrategy.LAYOUT, LoaderStrategy.OCR}),
    DocKind.IMAGE: frozenset({LoaderStrategy.OCR}),
}


@dataclass
class ExtractResult:
    text: str
    kind: DocKind
    strategy: LoaderStrategy
    tables_merged: int = 0
    warnings: list[str] = field(default_factory=list)


@dataclass
class _Table:
    name: str
    headers: list[str]
    rows: list[list[str]]


def decode_bytes(raw: bytes) -> str:
    """Decode text bytes with encodings common in mixed-language uploads."""
    for encoding in _TEXT_ENCODINGS:
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def detect_kind(filename: str, mime_type: str | None = None) -> DocKind:
    """Infer document kind from extension, then Content-Type."""
    ext = Path(filename).suffix.lower()
    if ext in _LEGACY_HINTS:
        raise ValueError(_LEGACY_HINTS[ext])
    if ext in _EXT_KIND:
        return _EXT_KIND[ext]
    mime = (mime_type or "").split(";", 1)[0].strip().lower()
    if mime in _MIME_KIND:
        return _MIME_KIND[mime]
    if mime.startswith("text/"):
        return DocKind.TEXT
    if mime.startswith("image/"):
        return DocKind.IMAGE
    raise ValueError(
        f"unsupported document type {filename!r}; "
        "use txt, md, json, csv, html, pdf, docx, xlsx, pptx, or common images"
    )


def choose_strategy(
    kind: DocKind,
    path: Path | None = None,
    *,
    force_strategy: LoaderStrategy | str | None = None,
) -> LoaderStrategy:
    """Pick extract strategy; ``force_strategy`` must be valid for ``kind``."""
    if force_strategy is not None:
        strategy = LoaderStrategy(force_strategy)
        allowed = _KIND_STRATEGIES[kind]
        if strategy not in allowed:
            raise ValueError(
                f"{strategy} loader does not accept {kind} files "
                f"(allowed: {', '.join(sorted(item.value for item in allowed))})"
            )
        return strategy
    if kind == DocKind.PDF and path is not None and _pdf_looks_scanned(path):
        return LoaderStrategy.OCR
    if kind == DocKind.IMAGE:
        return LoaderStrategy.OCR
    if kind in {DocKind.CSV, DocKind.XLSX}:
        return LoaderStrategy.TABLE
    if kind == DocKind.HTML:
        return LoaderStrategy.MARKUP
    if kind in {DocKind.PDF, DocKind.DOCX, DocKind.PPTX}:
        return LoaderStrategy.LAYOUT
    return LoaderStrategy.TEXT


def extract_path(
    path: Path,
    *,
    filename: str | None = None,
    mime_type: str | None = None,
    force_kind: DocKind | str | None = None,
    force_strategy: LoaderStrategy | str | None = None,
    merge_tables: bool = True,
    ocr_lang: str = "chi_sim+eng",
) -> ExtractResult:
    """Read ``path`` and return structured extract result."""
    name = filename or path.name
    kind = DocKind(force_kind) if force_kind else detect_kind(name, mime_type)
    if force_kind:
        detected = None
        try:
            detected = detect_kind(name, mime_type)
        except ValueError:
            detected = None
        if detected is not None and detected != kind:
            raise ValueError(f"{kind} expected but file looks like {detected}")
    strategy = choose_strategy(kind, path, force_strategy=force_strategy)
    return _run_strategy(
        path,
        kind=kind,
        strategy=strategy,
        merge_tables=merge_tables,
        ocr_lang=ocr_lang,
        name=name,
    )


def load_from_ingest_data(
    data: dict,
    *,
    force_kind: DocKind | str | None = None,
    force_strategy: LoaderStrategy | str | None = None,
    merge_tables: bool = True,
    ocr_lang: str = "chi_sim+eng",
) -> ExtractResult:
    """Prefer the stored file; fall back to already-extracted ``raw_text``."""
    path_value = data.get("file_path")
    filename = str(data.get("filename") or "")
    mime_type = data.get("mime_type")
    if path_value:
        path = Path(path_value)
        if not path.is_file():
            raise ValueError(f"uploaded file not found: {path}")
        return extract_path(
            path,
            filename=filename or path.name,
            mime_type=mime_type,
            force_kind=force_kind,
            force_strategy=force_strategy,
            merge_tables=merge_tables,
            ocr_lang=ocr_lang,
        )
    raw_text = data.get("raw_text")
    if raw_text:
        kind = DocKind(force_kind) if force_kind else DocKind.TEXT
        strategy = (
            LoaderStrategy(force_strategy) if force_strategy else LoaderStrategy.TEXT
        )
        if strategy not in {
            LoaderStrategy.TEXT,
            LoaderStrategy.MARKUP,
            LoaderStrategy.TABLE,
        }:
            raise ValueError(f"{strategy} loader needs file_path")
        text = str(raw_text)
        tables_merged = 0
        if strategy == LoaderStrategy.MARKUP:
            text = _html_text(text)
        elif strategy == LoaderStrategy.TABLE:
            text, tables_merged = _format_tables(
                [_parse_csv_table(text, name="inline")],
                merge=merge_tables,
            )
        elif kind == DocKind.JSON:
            try:
                text = _json_text(text)
            except json.JSONDecodeError:
                pass
        text = text.strip()
        if not text:
            raise ValueError("extracted no text from inline raw_text")
        return ExtractResult(
            text=text,
            kind=kind,
            strategy=strategy,
            tables_merged=tables_merged,
        )
    raise ValueError("loader needs file_path or raw_text")


def _run_strategy(
    path: Path,
    *,
    kind: DocKind,
    strategy: LoaderStrategy,
    merge_tables: bool,
    ocr_lang: str,
    name: str,
) -> ExtractResult:
    warnings: list[str] = []
    tables_merged = 0
    if strategy == LoaderStrategy.TEXT:
        text = _extract_text_strategy(path, kind)
    elif strategy == LoaderStrategy.MARKUP:
        text = _html_text(decode_bytes(path.read_bytes()))
    elif strategy == LoaderStrategy.LAYOUT:
        text = _extract_layout(path, kind)
        if kind == DocKind.PDF and _is_sparse_text(
            text, page_count=_pdf_page_count(path)
        ):
            warnings.append("pdf text layer is sparse; prefer ocr loader for scans")
    elif strategy == LoaderStrategy.TABLE:
        tables = _extract_tables(path, kind)
        text, tables_merged = _format_tables(tables, merge=merge_tables)
    elif strategy == LoaderStrategy.OCR:
        text = _ocr_path(path, kind=kind, lang=ocr_lang)
    else:
        raise ValueError(f"unknown loader strategy {strategy}")
    text = text.strip()
    if not text:
        raise ValueError(f"extracted no text from {name!r} via {strategy}")
    return ExtractResult(
        text=text,
        kind=kind,
        strategy=strategy,
        tables_merged=tables_merged,
        warnings=warnings,
    )


def _extract_text_strategy(path: Path, kind: DocKind) -> str:
    raw = decode_bytes(path.read_bytes())
    if kind == DocKind.JSON:
        return _json_text(raw)
    return raw


def _extract_layout(path: Path, kind: DocKind) -> str:
    if kind == DocKind.PDF:
        return _pdf_text(path)
    if kind == DocKind.DOCX:
        return _docx_prose(path)
    if kind == DocKind.PPTX:
        return _pptx_text(path)
    raise ValueError(f"layout strategy does not support {kind}")


def _extract_tables(path: Path, kind: DocKind) -> list[_Table]:
    if kind == DocKind.CSV:
        return [
            _parse_csv_table(
                decode_bytes(path.read_bytes()),
                name=path.stem,
                tsv=path.suffix.lower() == ".tsv",
            )
        ]
    if kind == DocKind.XLSX:
        return _xlsx_tables(path)
    if kind == DocKind.DOCX:
        return _docx_tables(path)
    raise ValueError(f"table strategy does not support {kind}")


def _html_text(raw: str) -> str:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(raw, "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    return soup.get_text("\n", strip=True)


def _json_text(raw: str) -> str:
    payload = json.loads(raw)
    return json.dumps(payload, ensure_ascii=False, indent=2)


def _parse_csv_table(raw: str, *, name: str, tsv: bool = False) -> _Table:
    reader = csv.reader(io.StringIO(raw), delimiter="\t" if tsv else ",")
    matrix = [
        [cell.strip() for cell in row] for row in reader if any(c.strip() for c in row)
    ]
    if not matrix:
        return _Table(name=name, headers=[], rows=[])
    return _Table(name=name, headers=matrix[0], rows=matrix[1:])


def _xlsx_tables(path: Path) -> list[_Table]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        tables: list[_Table] = []
        for sheet in workbook.worksheets:
            matrix: list[list[str]] = []
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if cell is None else str(cell).strip() for cell in row]
                if any(cells):
                    matrix.append(cells)
            if not matrix:
                continue
            tables.append(_Table(name=sheet.title, headers=matrix[0], rows=matrix[1:]))
        return tables
    finally:
        workbook.close()


def _docx_tables(path: Path) -> list[_Table]:
    from docx import Document as DocxDocument

    document = DocxDocument(str(path))
    tables: list[_Table] = []
    for index, table in enumerate(document.tables, start=1):
        matrix: list[list[str]] = []
        for row in table.rows:
            cells: list[str] = []
            for cell in row.cells:
                value = cell.text.strip()
                # python-docx repeats merged-cell text; collapse adjacent duplicates.
                if cells and cells[-1] == value:
                    continue
                cells.append(value)
            if any(cells):
                matrix.append(cells)
        if not matrix:
            continue
        tables.append(
            _Table(name=f"table-{index}", headers=matrix[0], rows=matrix[1:])
        )
    return tables


def _header_key(headers: list[str]) -> tuple[str, ...]:
    return tuple(re.sub(r"\s+", " ", header).strip().lower() for header in headers)


def _format_tables(tables: list[_Table], *, merge: bool) -> tuple[str, int]:
    usable = [table for table in tables if table.headers or table.rows]
    if not usable:
        return "", 0
    if not merge:
        return "\n\n".join(_render_table(table) for table in usable), 0

    groups: dict[tuple[str, ...], _Table] = {}
    order: list[tuple[str, ...]] = []
    merged = 0
    for table in usable:
        key = _header_key(table.headers)
        if key in groups and key != ():
            target = groups[key]
            target.rows.extend(table.rows)
            if table.name not in target.name:
                target.name = f"{target.name}+{table.name}"
            merged += 1
            continue
        groups[key] = _Table(
            name=table.name,
            headers=list(table.headers),
            rows=[list(row) for row in table.rows],
        )
        order.append(key)
    return "\n\n".join(_render_table(groups[key]) for key in order), merged


def _render_table(table: _Table) -> str:
    lines = [f"# {table.name}"]
    if table.headers:
        lines.append("\t".join(table.headers))
    for row in table.rows:
        width = len(table.headers) if table.headers else len(row)
        cells = list(row[:width]) + [""] * max(0, width - len(row))
        lines.append("\t".join(cells))
    return "\n".join(lines)


def _pdf_text(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = [(page.extract_text() or "").strip() for page in reader.pages]
    return "\n\n".join(part for part in pages if part)


def _pdf_page_count(path: Path) -> int:
    from pypdf import PdfReader

    return max(1, len(PdfReader(str(path)).pages))


def _pdf_looks_scanned(path: Path) -> bool:
    try:
        text = _pdf_text(path)
        pages = _pdf_page_count(path)
    except Exception:
        return True
    return _is_sparse_text(text, page_count=pages)


def _is_sparse_text(text: str, *, page_count: int) -> bool:
    chars = len(re.sub(r"\s+", "", text or ""))
    return chars < _PDF_SPARSE_CHARS_PER_PAGE * max(1, page_count)


def _docx_prose(path: Path) -> str:
    from docx import Document as DocxDocument

    document = DocxDocument(str(path))
    parts = [para.text.strip() for para in document.paragraphs if para.text.strip()]
    if document.tables:
        parts.append(
            f"[docx contains {len(document.tables)} table(s); use table loader to merge]"
        )
    return "\n".join(parts)


def _pptx_text(path: Path) -> str:
    from pptx import Presentation

    presentation = Presentation(str(path))
    parts: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        parts.append(f"# Slide {index}")
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            text = shape.text_frame.text.strip()
            if text:
                parts.append(text)
    return "\n".join(parts)


def _ocr_path(path: Path, *, kind: DocKind, lang: str) -> str:
    try:
        from PIL import Image
    except ImportError as exc:
        raise ValueError(
            "OCR requires Pillow; install with: pip install 'raglab[ocr]'"
        ) from exc
    try:
        import pytesseract
    except ImportError as exc:
        raise ValueError(
            "OCR requires pytesseract; install with: pip install 'raglab[ocr]' "
            "and a system Tesseract binary"
        ) from exc

    if kind == DocKind.IMAGE:
        return pytesseract.image_to_string(Image.open(path), lang=lang)

    if kind == DocKind.PDF:
        try:
            from pdf2image import convert_from_path
        except ImportError as exc:
            raise ValueError(
                "scanned PDF OCR requires pdf2image (and poppler); "
                "install with: pip install 'raglab[ocr]'"
            ) from exc
        images = convert_from_path(str(path))
        parts: list[str] = []
        for index, image in enumerate(images, start=1):
            page_text = pytesseract.image_to_string(image, lang=lang).strip()
            if page_text:
                parts.append(f"# Page {index}\n{page_text}")
        return "\n\n".join(parts)

    raise ValueError(f"ocr strategy does not support {kind}")
