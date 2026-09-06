"""End-to-end pipelines: every builtin plugin appears in at least one full run."""

from __future__ import annotations

import math
from pathlib import Path
from typing import Any
from uuid import UUID, uuid4

import pytest
from docx import Document as DocxDocument

from engine.retrieval import cosine_similarity, tokenize
from engine.runner import ordered_calls, run_pipeline
from models.enums import PipelineKind
from plugins.builtin import ALL_BUILTIN_PLUGINS, register_builtin
from plugins.context import MemoryTrace
from plugins.registry import PluginRegistry

POLICY_TEXT = (
    "# Refund\n\n"
    "Customers may refund Widget X within 7 days of purchase.\n\n"
    "# Shipping\n\n"
    "Orders ship within three business days to mainland addresses.\n"
)
QUERY = "How long can I refund Widget X?"
GOLD_QUOTE = "refund Widget X within 7 days"


class InMemoryKbRepo:
    """Stores chunks/embeddings from ingest and serves dense/BM25 retrieval."""

    def __init__(self) -> None:
        self.chunks: dict[UUID, dict[str, Any]] = {}
        self.embeddings: list[dict[str, Any]] = []
        self.collection_dim: int | None = None

    async def insert_chunk(self, **kwargs: Any) -> dict[str, Any]:
        chunk_id = uuid4()
        row = {"id": chunk_id, **kwargs}
        self.chunks[chunk_id] = row
        return {"id": chunk_id, "ordinal": kwargs["ordinal"], "content": kwargs["content"]}

    async def insert_embedding(self, **kwargs: Any) -> None:
        self.embeddings.append(kwargs)

    async def update_collection_dim(self, collection_id: Any, dim: int) -> None:
        self.collection_dim = dim

    def _hit(self, chunk_id: UUID, score: float | None = None) -> dict[str, Any]:
        chunk = self.chunks[chunk_id]
        parent_id = chunk.get("parent_chunk_id")
        parent = self.chunks.get(parent_id) if parent_id else None
        expanded = parent["content"] if parent else chunk["content"]
        row: dict[str, Any] = {
            "chunk_id": chunk_id,
            "content": chunk["content"],
            "expanded_content": expanded,
            "metadata": chunk.get("metadata") or {},
            "parent_chunk_id": parent_id,
        }
        if score is not None:
            row["score"] = score
        return row

    async def dense_search(
        self,
        *,
        vector_collection_id: UUID,
        query_embedding: list[float],
        top_k: int,
        score_threshold: float | None,
    ) -> list[dict[str, Any]]:
        scored: list[tuple[float, UUID]] = []
        for item in self.embeddings:
            if item["vector_collection_id"] != vector_collection_id:
                continue
            score = cosine_similarity(query_embedding, item["embedding"])
            if score_threshold is not None and score < score_threshold:
                continue
            scored.append((score, item["chunk_id"]))
        scored.sort(key=lambda pair: pair[0], reverse=True)
        return [self._hit(chunk_id, score) for score, chunk_id in scored[:top_k]]

    async def list_collection_chunks(self, vector_collection_id: UUID) -> list[dict[str, Any]]:
        ids = {
            item["chunk_id"]
            for item in self.embeddings
            if item["vector_collection_id"] == vector_collection_id
        }
        return [self._hit(chunk_id) for chunk_id in ids]


class PipelineCtx:
    """Duck-typed PluginContext: real stage order, fake models + KB."""

    def __init__(self, kb_repo: InMemoryKbRepo) -> None:
        self.kb_repo = kb_repo
        self.trace = MemoryTrace()
        self.default_embedder_binding_id = uuid4()
        self.default_generator_binding_id = uuid4()
        self.extra: dict[str, Any] = {}
        self.chat_calls = 0
        self.embed_calls = 0

    async def embed_texts(self, texts: list[str], binding_id: Any) -> list[list[float]]:
        self.embed_calls += 1
        return [_bow_embedding(text) for text in texts]

    async def chat_complete(
        self,
        messages: list[dict[str, str]],
        binding_id: Any,
        *,
        temperature: float,
        max_tokens: int,
    ) -> str:
        self.chat_calls += 1
        system = (messages[0].get("content") or "") if messages else ""
        user = ""
        for msg in reversed(messages):
            if msg.get("role") == "user":
                user = msg.get("content") or ""
                break
        lower_system = system.lower()

        if "rewrite the latest user question" in lower_system:
            return "Widget X refund policy within 7 days"

        if "hypothetical answer" in lower_system:
            return "Customers may refund Widget X within seven days of purchase."

        if "diverse search queries" in lower_system:
            return (
                "Widget X refund policy\n"
                "Widget X return window 7 days\n"
                "how to refund Widget X"
            )

        if "score how relevant" in lower_system or "judge if the passage" in lower_system:
            score = 0.92 if "refund" in user.lower() else 0.15
            return f'{{"score": {score}}}'

        if "grounded in the passages" in lower_system:
            return '{"score": 0.9}'

        if "answer using only the retrieved passages" in lower_system:
            return "You may refund Widget X within 7 days of purchase. [1]"

        return "ok"

    def token_callback(self):
        return None


def _bow_embedding(text: str, dim: int = 48) -> list[float]:
    vec = [0.0] * dim
    for token in tokenize(text):
        vec[hash(token) % dim] += 1.0
    norm = math.sqrt(sum(v * v for v in vec))
    if norm == 0:
        return vec
    return [v / norm for v in vec]


def _registry() -> PluginRegistry:
    registry = PluginRegistry()
    register_builtin(registry)
    return registry


def _slot(stage: str, name: str, params: dict[str, Any] | None = None) -> dict[str, Any]:
    return {"stage": stage, "bindings": [{"name": name, "params": params or {}}]}


def _ingest_slots(
    *,
    loader: str = "auto",
    loader_params: dict[str, Any] | None = None,
    chunker: str = "recursive",
    chunker_params: dict[str, Any] | None = None,
    embedder: str = "openai_embedder",
    embedder_params: dict[str, Any] | None = None,
    indexer: str = "pgvector",
) -> list[dict[str, Any]]:
    default_chunker = {
        "recursive": {"chunk_size": 80, "overlap": 8},
        "semantic": {"similarity_threshold": 0.5, "max_tokens": 120},
        "parent_child": {"parent_max_tokens": 160, "child_max_tokens": 60, "overlap": 0},
        "heading": {"chunk_size": 200, "overlap": 0},
    }
    return [
        _slot("loader", loader, loader_params),
        _slot("chunker", chunker, chunker_params or default_chunker.get(chunker)),
        _slot("embedder", embedder, embedder_params or {"batch_size": 8}),
        _slot("indexer", indexer, {"metric": "cosine"}),
    ]


def _query_slots(
    *,
    transformer: str = "rewrite",
    transformer_params: dict[str, Any] | None = None,
    retrievers: list[tuple[str, dict[str, Any]]] | None = None,
    fusion: str = "rrf",
    reranker: str = "bge-reranker",
    reranker_params: dict[str, Any] | None = None,
    grader: str = "crag",
    grader_params: dict[str, Any] | None = None,
    compressor: str = "top_n",
    compressor_params: dict[str, Any] | None = None,
    generator: str = "chat",
) -> list[dict[str, Any]]:
    retrievers = retrievers or [("dense", {"top_k": 5}), ("bm25", {"top_k": 5})]
    default_transformer = {
        "passthrough": {},
        "rewrite": {"history_turns": 0},
        "hyde": {"n_hypothetical": 1, "history_turns": 0},
        "multi_query": {"n_queries": 3, "history_turns": 0},
    }
    retriever_slot: dict[str, Any]
    if len(retrievers) == 1:
        name, params = retrievers[0]
        retriever_slot = _slot("retriever", name, params)
    else:
        retriever_slot = {
            "stage": "retriever",
            "mode": "ensemble",
            "bindings": [{"name": name, "params": params} for name, params in retrievers],
        }
    return [
        _slot(
            "query_transformer",
            transformer,
            transformer_params or default_transformer.get(transformer),
        ),
        retriever_slot,
        _slot("fusion", fusion, {"k": 60}),
        _slot("reranker", reranker, reranker_params or {"top_n": 3}),
        _slot("grader", grader, grader_params or {"accept_threshold": 0.5}),
        _slot("compressor", compressor, compressor_params or {"n": 2}),
        _slot("generator", generator, {"temperature": 0.0, "max_tokens": 128}),
    ]


def _make_loader_fixture(tmp_path: Path, kind: str) -> dict[str, Any]:
    """Build ingest input suited to the loader under test."""
    if kind in {"auto", "text"}:
        path = tmp_path / "policy.md"
        path.write_text(POLICY_TEXT, encoding="utf-8")
        return {"file_path": str(path), "filename": path.name, "raw_text": ""}

    if kind == "markup":
        path = tmp_path / "policy.html"
        path.write_text(
            "<html><body><h1>Refund</h1>"
            "<p>Customers may refund Widget X within 7 days of purchase.</p>"
            "<script>ignore()</script></body></html>",
            encoding="utf-8",
        )
        return {"file_path": str(path), "filename": path.name, "raw_text": ""}

    if kind == "layout":
        path = tmp_path / "policy.docx"
        document = DocxDocument()
        document.add_heading("Refund", level=1)
        document.add_paragraph("Customers may refund Widget X within 7 days of purchase.")
        document.add_heading("Shipping", level=1)
        document.add_paragraph("Orders ship within three business days.")
        document.save(path)
        return {"file_path": str(path), "filename": path.name, "raw_text": ""}

    if kind == "table":
        path = tmp_path / "policy.csv"
        path.write_text(
            "topic,detail\n"
            "Refund,Customers may refund Widget X within 7 days of purchase\n"
            "Shipping,Orders ship within three business days\n",
            encoding="utf-8",
        )
        return {"file_path": str(path), "filename": path.name, "raw_text": ""}

    if kind == "ocr":
        try:
            from PIL import Image, ImageDraw
        except ImportError:
            pytest.skip("raglab[ocr] / Pillow not installed")
        path = tmp_path / "policy.png"
        image = Image.new("RGB", (640, 160), color="white")
        draw = ImageDraw.Draw(image)
        draw.text((12, 40), "Customers may refund Widget X within 7 days", fill="black")
        image.save(path)
        return {"file_path": str(path), "filename": path.name, "raw_text": ""}

    raise AssertionError(f"unknown loader fixture kind: {kind}")


async def _run_ingest_query(
    tmp_path: Path,
    *,
    loader: str = "auto",
    loader_params: dict[str, Any] | None = None,
    chunker: str = "recursive",
    chunker_params: dict[str, Any] | None = None,
    embedder: str = "openai_embedder",
    transformer: str = "rewrite",
    retrievers: list[tuple[str, dict[str, Any]]] | None = None,
    reranker: str = "bge-reranker",
    compressor: str = "top_n",
    grader_threshold: float = 0.5,
    fixture_kind: str | None = None,
    run_evaluators: bool = False,
) -> tuple[dict[str, Any], dict[str, Any], set[tuple[str, str]]]:
    """Run one full ingest→query recipe; return results and plugins used."""
    registry = _registry()
    kb = InMemoryKbRepo()
    ctx = PipelineCtx(kb)
    collection_id = uuid4()
    fixture = _make_loader_fixture(tmp_path, fixture_kind or loader)

    ingest_slots = _ingest_slots(
        loader=loader,
        loader_params=loader_params,
        chunker=chunker,
        chunker_params=chunker_params,
        embedder=embedder,
    )
    query_slots = _query_slots(
        transformer=transformer,
        retrievers=retrievers,
        reranker=reranker,
        grader_params={"accept_threshold": grader_threshold},
        compressor=compressor,
    )

    used: set[tuple[str, str]] = set()
    for slot in ingest_slots + query_slots:
        for binding in slot["bindings"]:
            used.add((slot["stage"], binding["name"]))

    try:
        ingest = await run_pipeline(
            calls=ordered_calls(ingest_slots, PipelineKind.INGEST),
            data={
                **fixture,
                "document_version_id": uuid4(),
                "vector_collection_id": collection_id,
            },
            ctx=ctx,  # type: ignore[arg-type]
            registry=registry,
        )
    except Exception as exc:
        if loader == "ocr":
            pytest.skip(f"OCR runtime unavailable: {exc}")
        raise

    assert ingest.get("raw_text"), f"{loader}/{chunker} produced empty raw_text"
    assert "Widget" in ingest["raw_text"] or "refund" in ingest["raw_text"].lower()
    assert ingest["chunks"], f"chunker {chunker} produced no chunks"
    assert ingest["indexed_count"] >= 1
    assert ingest["chunker_plugin"] == chunker
    assert ingest["embedder_plugin"] == embedder

    ctx.trace = MemoryTrace()
    query = await run_pipeline(
        calls=ordered_calls(query_slots, PipelineKind.QUERY),
        data={
            "query": QUERY,
            "history": [],
            "vector_collection_id": collection_id,
        },
        ctx=ctx,  # type: ignore[arg-type]
        registry=registry,
    )

    assert query.get("rewritten_query")
    assert query.get("retrieved"), f"no passages after {transformer}/{retrievers}/{reranker}"
    assert query.get("answer")
    assert "[1]" in query["answer"] or "7 days" in query["answer"]
    assert [s.stage for s in ctx.trace.spans] == [
        "query_transformer",
        "retriever",
        "fusion",
        "reranker",
        "grader",
        "compressor",
        "generator",
    ]

    if run_evaluators:
        query["gold_quotes"] = [GOLD_QUOTE]
        for name in ("recall_at_k", "mrr", "faithfulness"):
            plugin = registry.get("evaluator", name)
            assert plugin is not None
            query = await plugin.execute(query, {"k": 5} if name == "recall_at_k" else {}, ctx)
            used.add(("evaluator", name))
        assert query["metrics"]["recall_at_k"] is not None
        assert query["metrics"]["mrr"] is not None
        assert query["metrics"]["faithfulness"] == 0.9

    return ingest, query, used


# --- one full pipeline per plugin variant ---


@pytest.mark.parametrize(
    ("loader", "loader_params"),
    [
        ("auto", None),
        ("text", None),
        ("markup", None),
        ("layout", None),
        ("table", {"merge_tables": True}),
        ("ocr", {"ocr_lang": "eng"}),
    ],
)
async def test_pipeline_each_loader(
    tmp_path: Path, loader: str, loader_params: dict[str, Any] | None
) -> None:
    await _run_ingest_query(tmp_path, loader=loader, loader_params=loader_params)


@pytest.mark.parametrize("chunker", ["recursive", "semantic", "parent_child", "heading"])
async def test_pipeline_each_chunker(tmp_path: Path, chunker: str) -> None:
    await _run_ingest_query(tmp_path, loader="auto", chunker=chunker, fixture_kind="auto")


@pytest.mark.parametrize("embedder", ["openai_embedder", "local_embedder"])
async def test_pipeline_each_embedder(tmp_path: Path, embedder: str) -> None:
    await _run_ingest_query(tmp_path, embedder=embedder)


@pytest.mark.parametrize(
    "transformer", ["passthrough", "rewrite", "hyde", "multi_query"]
)
async def test_pipeline_each_query_transformer(tmp_path: Path, transformer: str) -> None:
    ingest, query, _ = await _run_ingest_query(tmp_path, transformer=transformer)
    if transformer == "passthrough":
        assert query["rewritten_query"] == QUERY
    elif transformer == "hyde":
        assert query.get("hypothetical_docs")
        assert query.get("search_queries")
    elif transformer == "multi_query":
        assert len(query.get("search_queries") or []) >= 1
    assert ingest["indexed_count"] >= 1


@pytest.mark.parametrize(
    "retrievers",
    [
        [("dense", {"top_k": 5})],
        [("bm25", {"top_k": 5})],
        [("dense", {"top_k": 5}), ("bm25", {"top_k": 5})],
    ],
    ids=["dense", "bm25", "ensemble"],
)
async def test_pipeline_each_retriever_mode(
    tmp_path: Path, retrievers: list[tuple[str, dict[str, Any]]]
) -> None:
    _, query, _ = await _run_ingest_query(tmp_path, retrievers=retrievers)
    for name, _params in retrievers:
        assert name in query["retrieved_lists"]
        assert query["retrieved_lists"][name]


@pytest.mark.parametrize("reranker", ["none", "bge-reranker"])
async def test_pipeline_each_reranker(tmp_path: Path, reranker: str) -> None:
    await _run_ingest_query(tmp_path, reranker=reranker, grader_threshold=0.0)


@pytest.mark.parametrize(
    ("compressor", "compressor_params"),
    [("none", {}), ("top_n", {"n": 2})],
)
async def test_pipeline_each_compressor(
    tmp_path: Path, compressor: str, compressor_params: dict[str, Any]
) -> None:
    _, query, _ = await _run_ingest_query(
        tmp_path,
        compressor=compressor,
        grader_threshold=0.0,
    )
    if compressor == "top_n":
        assert len(query["retrieved"]) <= 2


async def test_pipeline_grader_and_generator_and_fusion(tmp_path: Path) -> None:
    """crag + chat + rrf are singletons; still run a dedicated full recipe."""
    _, query, used = await _run_ingest_query(
        tmp_path,
        transformer="rewrite",
        retrievers=[("dense", {"top_k": 5}), ("bm25", {"top_k": 5})],
        reranker="bge-reranker",
        compressor="top_n",
    )
    assert ("grader", "crag") in used
    assert ("generator", "chat") in used
    assert ("fusion", "rrf") in used
    assert ("indexer", "pgvector") in used
    assert query["crag_action"] == "correct"


async def test_pipeline_with_all_evaluators(tmp_path: Path) -> None:
    await _run_ingest_query(tmp_path, run_evaluators=True)


def _ocr_deps_available() -> bool:
    try:
        import pytesseract  # noqa: F401
        from PIL import Image  # noqa: F401
    except ImportError:
        return False
    return True


async def test_every_builtin_plugin_covered_by_a_full_pipeline(tmp_path: Path) -> None:
    """Sweep recipes until every registered builtin has been in a run_pipeline path."""
    recipes: list[dict[str, Any]] = [
        {"loader": "auto"},
        {"loader": "text", "fixture_kind": "text"},
        {"loader": "markup", "fixture_kind": "markup"},
        {"loader": "layout", "fixture_kind": "layout"},
        {"loader": "table", "loader_params": {"merge_tables": True}, "fixture_kind": "table"},
        {"chunker": "recursive"},
        {"chunker": "semantic"},
        {"chunker": "parent_child"},
        {"chunker": "heading"},
        {"embedder": "openai_embedder"},
        {"embedder": "local_embedder"},
        {"transformer": "passthrough"},
        {"transformer": "rewrite"},
        {"transformer": "hyde"},
        {"transformer": "multi_query"},
        {"retrievers": [("dense", {"top_k": 5})]},
        {"retrievers": [("bm25", {"top_k": 5})]},
        {"retrievers": [("dense", {"top_k": 5}), ("bm25", {"top_k": 5})]},
        {"reranker": "none", "grader_threshold": 0.0},
        {"reranker": "bge-reranker"},
        {"compressor": "none", "grader_threshold": 0.0},
        {"compressor": "top_n"},
        {"run_evaluators": True},
    ]
    if _ocr_deps_available():
        recipes.append(
            {"loader": "ocr", "loader_params": {"ocr_lang": "eng"}, "fixture_kind": "ocr"}
        )

    covered: set[tuple[str, str]] = set()
    for index, recipe in enumerate(recipes):
        case_dir = tmp_path / f"case_{index}"
        case_dir.mkdir()
        _ingest, _query, used = await _run_ingest_query(case_dir, **recipe)
        covered |= used

    expected = {(p.stage.value, p.name) for p in ALL_BUILTIN_PLUGINS}
    if not _ocr_deps_available():
        expected.discard(("loader", "ocr"))
    missing = sorted(expected - covered)
    assert missing == [], f"plugins never run in a full pipeline: {missing}"