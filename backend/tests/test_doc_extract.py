from pathlib import Path

from docx import Document as DocxDocument
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from engine.doc_extract import (
    DocKind,
    LoaderStrategy,
    choose_strategy,
    detect_kind,
    extract_path,
    load_from_ingest_data,
)
from plugins.builtin.loader_stage import auto, table, text


def test_detect_kind_from_extension_and_mime() -> None:
    assert detect_kind("notes.md") is DocKind.TEXT
    assert detect_kind("report.PDF") is DocKind.PDF
    assert detect_kind("upload.bin", "application/pdf") is DocKind.PDF
    assert detect_kind("scan.png") is DocKind.IMAGE


def test_detect_kind_rejects_legacy_office() -> None:
    try:
        detect_kind("old.doc")
        raise AssertionError("expected ValueError")
    except ValueError as exc:
        assert "docx" in str(exc)


def test_choose_strategy_routes_by_kind() -> None:
    assert choose_strategy(DocKind.HTML) is LoaderStrategy.MARKUP
    assert choose_strategy(DocKind.CSV) is LoaderStrategy.TABLE
    assert choose_strategy(DocKind.IMAGE) is LoaderStrategy.OCR
    assert choose_strategy(DocKind.DOCX) is LoaderStrategy.LAYOUT


def test_extract_text_html_json_csv(tmp_path: Path) -> None:
    (tmp_path / "a.txt").write_text("hello lab", encoding="utf-8")
    (tmp_path / "a.html").write_text(
        "<html><body><p>Hi</p><script>secret()</script></body></html>",
        encoding="utf-8",
    )
    (tmp_path / "a.json").write_text('{"k": "v"}', encoding="utf-8")
    (tmp_path / "a.csv").write_text("a,b\n1,2\n", encoding="utf-8")

    assert extract_path(tmp_path / "a.txt").text == "hello lab"
    html = extract_path(tmp_path / "a.html")
    assert html.kind is DocKind.HTML
    assert html.strategy is LoaderStrategy.MARKUP
    assert "Hi" in html.text
    assert "secret" not in html.text
    assert '"k"' in extract_path(tmp_path / "a.json").text
    csv_result = extract_path(tmp_path / "a.csv")
    assert csv_result.strategy is LoaderStrategy.TABLE
    assert "1\t2" in csv_result.text


def test_table_loader_merges_same_header_sheets(tmp_path: Path) -> None:
    path = tmp_path / "m.xlsx"
    workbook = Workbook()
    first = workbook.active
    first.title = "s1"
    first["A1"] = "a"
    first["B1"] = "b"
    first["A2"] = "1"
    first["B2"] = "2"
    second = workbook.create_sheet("s2")
    second["A1"] = "a"
    second["B1"] = "b"
    second["A2"] = "3"
    second["B2"] = "4"
    workbook.save(path)

    merged = extract_path(path, merge_tables=True)
    assert merged.tables_merged == 1
    assert "1" in merged.text and "3" in merged.text

    separate = extract_path(path, merge_tables=False)
    assert separate.tables_merged == 0
    assert separate.text.count("# ") >= 2


def test_extract_office_layout(tmp_path: Path) -> None:
    docx_path = tmp_path / "n.docx"
    document = DocxDocument()
    document.add_paragraph("Word body")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "c1"
    table.cell(0, 1).text = "c2"
    document.save(docx_path)

    pptx_path = tmp_path / "n.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "Slide body"
    presentation.save(pptx_path)

    docx = extract_path(docx_path)
    assert docx.strategy is LoaderStrategy.LAYOUT
    assert "Word body" in docx.text

    table_only = extract_path(docx_path, force_strategy=LoaderStrategy.TABLE)
    assert "c1" in table_only.text

    assert "Slide body" in extract_path(pptx_path).text


async def test_auto_and_specialized_loaders(tmp_path: Path) -> None:
    path = tmp_path / "note.md"
    path.write_text("alpha", encoding="utf-8")
    result = await auto.execute(
        {"file_path": str(path), "filename": path.name, "raw_text": ""},
        {},
        None,
    )
    assert result["raw_text"] == "alpha"
    assert result["loader_kind"] == "text"
    assert result["loader_strategy"] == "text"

    html = tmp_path / "x.html"
    html.write_text("<p>Hi</p>", encoding="utf-8")
    try:
        await text.execute(
            {"file_path": str(html), "filename": html.name, "raw_text": ""},
            {},
            None,
        )
        raise AssertionError("expected strategy mismatch")
    except ValueError as exc:
        assert "html" in str(exc)


async def test_table_plugin_merge_param(tmp_path: Path) -> None:
    path = tmp_path / "m.xlsx"
    workbook = Workbook()
    first = workbook.active
    first.title = "s1"
    first["A1"] = "h"
    first["A2"] = "1"
    second = workbook.create_sheet("s2")
    second["A1"] = "h"
    second["A2"] = "2"
    workbook.save(path)
    result = await table.execute(
        {"file_path": str(path), "filename": path.name, "raw_text": ""},
        {"merge_tables": True},
        None,
    )
    assert result["tables_merged"] == 1
    assert "1" in result["raw_text"] and "2" in result["raw_text"]


async def test_auto_accepts_inline_raw_text() -> None:
    result = await auto.execute({"raw_text": "already extracted"}, {}, None)
    assert result["raw_text"] == "already extracted"


def test_load_from_ingest_prefers_file(tmp_path: Path) -> None:
    path = tmp_path / "a.txt"
    path.write_text("from-disk", encoding="utf-8")
    result = load_from_ingest_data(
        {"file_path": str(path), "filename": "a.txt", "raw_text": "stale"}
    )
    assert result.text == "from-disk"
    assert result.kind is DocKind.TEXT
